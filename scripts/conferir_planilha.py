#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
conferir_planilha.py
====================
Confere um .xlsx já gerado contra o contrato de formato descrito em
references/formato-planilha.md, e devolve as violações como lista.

Por que existe: até setembro/2026 a skill produzia e não conferia. O contrato
estava escrito, o código o implementava, e ninguém comparava o resultado com
ele. Foi assim que a logo ficou achatada em todo cliente por meses (razão 2,40
em vez de 3,56), que o recalc passou a cortar a imagem sem avisar, que a
planilha começou a abrir em meia tela, e que o nome do arquivo regrediu de
com hifen para o nome colado. Todas essas descobertas foram do
Daniel, abrindo o arquivo e olhando. São verificações de um segundo cada.

As constantes vêm de build_relatorio.py e recalc.py de propósito, por import e
não por cópia: um conferidor com os próprios números vira uma terceira versão
da verdade e passa a discordar dos outros dois em silêncio.

Uso:
  python3 conferir_planilha.py ARQUIVO.xlsx
  python3 conferir_planilha.py ARQUIVO.xlsx --itens 15 --total 977.67 \\
      --nome-entrega "{Cliente}"

Saída: JSON com "ok", "falhas" (só o texto do que reprovou) e "verificacoes"
(todas, no formato text/passed/evidence). Código de saída 0 se tudo passou,
2 se algo reprovou, para o orquestrador poder ramificar sem parsear texto.
"""
import argparse
import json
import os
import re
import sys
import zipfile

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)

from build_relatorio import WIDTHS, LOGO_ANCORA  # noqa: E402
from recalc import JANELA  # noqa: E402

# 566x159 é o PNG original da logo. A razão é o que não pode mudar; a largura,
# sim, se a coluna B mudar. Ver formato-planilha.md > Injeção da logo.
RAZAO_LOGO = 3.5597
TOLERANCIA_RAZAO = 0.02


def conferir(caminho, itens=None, total=None, nome_entrega=None):
    from openpyxl import load_workbook

    checks = []

    def add(texto, passou, evidencia):
        checks.append({'text': texto, 'passed': bool(passou), 'evidence': str(evidencia)})

    if not os.path.isfile(caminho):
        add('O arquivo existe', False, caminho)
        return checks

    nome = os.path.basename(caminho)
    wb = load_workbook(caminho)
    ws = wb.active
    wsv = load_workbook(caminho, data_only=True).active
    z = zipfile.ZipFile(caminho)
    partes = z.namelist()

    # ---- nome do arquivo -------------------------------------------------
    add('O nome do arquivo não tem espaço', ' ' not in nome, nome)
    if nome_entrega:
        prefixo = nome_entrega.replace(' ', '-')
        add('O nome começa com "%s", com espaço virando hífen' % prefixo,
            nome.startswith(prefixo), nome)

    # ---- logo ------------------------------------------------------------
    midia = [n for n in partes if n.startswith('xl/media/')]
    add('A logo está embutida no arquivo', bool(midia), midia or 'sem xl/media')
    add('O drawing da logo sobreviveu ao pipeline',
        'xl/drawings/drawing1.xml' in partes,
        'drawing1.xml presente' if 'xl/drawings/drawing1.xml' in partes else 'ausente')

    if 'xl/drawings/drawing1.xml' in partes:
        xml = z.read('xl/drawings/drawing1.xml').decode('utf-8')
        m = re.search(r'<a:ext cx="(\d+)" cy="(\d+)"/>', xml)
        if m:
            cx, cy = int(m.group(1)), int(m.group(2))
            razao = cx / float(cy)
            add('A logo mantém a proporção do PNG original (%.4f)' % RAZAO_LOGO,
                abs(razao - RAZAO_LOGO) < TOLERANCIA_RAZAO,
                'cx=%d cy=%d razão=%.4f' % (cx, cy, razao))
        else:
            add('A logo mantém a proporção do PNG original', False,
                '<a:ext> não encontrado no drawing1.xml')
        m = re.search(r'<xdr:to>.*?<xdr:colOff>(\d+)</xdr:colOff>', xml, re.S)
        if m:
            fim = int(m.group(1))
            add('A âncora da logo cabe na coluna B, então o recalc não a corta',
                fim <= LOGO_ANCORA['to_col_off'],
                'termina em %d, limite %d' % (fim, LOGO_ANCORA['to_col_off']))

    # ---- geometria da folha ---------------------------------------------
    alt2 = ws.row_dimensions[2].height
    add('A linha 2 tem 66 pt, senão a logo não cabe', alt2 == 66.0, alt2)
    alt3 = ws.row_dimensions[3].height
    add('A linha 3 tem 42 pt', alt3 == 42.0, alt3)

    fora = []
    for col, largura in WIDTHS.items():
        atual = ws.column_dimensions[col].width
        if atual is None or abs(atual - largura) > 0.5:
            if col == 'H' and not ws.cell(3, 8).value:
                continue  # cliente sem coluna extra
            fora.append('%s=%s (esperado %s)' % (col, atual, largura))
    add('As larguras de coluna batem com o contrato', not fora, fora or 'todas ok')

    # ---- conteúdo --------------------------------------------------------
    linhas = [i for i in range(4, ws.max_row + 1) if ws.cell(i, 4).value is not None]
    add('A planilha tem pelo menos uma demanda', bool(linhas), '%d demandas' % len(linhas))
    if itens is not None:
        add('A planilha tem exatamente %d demandas' % itens,
            len(linhas) == itens, '%d demandas' % len(linhas))

    if linhas:
        soma = round(sum(float(ws.cell(i, 4).value) for i in linhas), 2)
        if total is not None:
            add('O total soma R$ %.2f' % total, abs(soma - total) < 0.01, 'soma=%.2f' % soma)

        bs = [ws.cell(i, 2).value for i in linhas]
        vazias = [i for i, v in zip(linhas, bs) if not str(v or '').strip()]
        numericas = [v for v in bs if str(v or '').strip().isdigit()]
        add('Coluna B com Task Name em texto, sem vazia e sem numeração sequencial',
            not vazias and not numericas,
            'vazias=%s numéricas=%s exemplos=%s' % (vazias, numericas, bs[:2]))

        datas = [ws.cell(i, 6).value for i in linhas]
        ordenado = all(datas[i] is None or datas[i + 1] is None or datas[i] <= datas[i + 1]
                       for i in range(len(datas) - 1))
        add('Demandas em ordem cronológica', ordenado,
            ' -> '.join(d.strftime('%d/%m') if hasattr(d, 'strftime') else str(d)
                        for d in datas[:6]))

        revisar = [ws.cell(i, 8).value for i in linhas
                   if str(ws.cell(i, 8).value or '').startswith('REVISAR')]
        add('Nenhuma célula da coluna extra ficou em REVISAR',
            not revisar, revisar or 'nenhuma')

    cab = ws.cell(3, 2)
    add('Cabeçalho da linha 3 em Calibri 14',
        cab.font.name == 'Calibri' and float(cab.font.size or 0) == 14.0,
        '%s %s' % (cab.font.name, cab.font.size))
    if linhas:
        dado = ws.cell(linhas[0], 3)
        add('Linhas de dados em Calibri 12',
            dado.font.name == 'Calibri' and float(dado.font.size or 0) == 12.0,
            '%s %s' % (dado.font.name, dado.font.size))

    # ---- total e cache do SUM -------------------------------------------
    cols = [c for c in 'BCDEFGH' if ws.cell(3, 'BCDEFGH'.index(c) + 2).value]
    ultima = cols[-1] if cols else 'G'
    formula = ws['%s2' % ultima].value
    cache = wsv['%s2' % ultima].value
    add('Total em fórmula SUM na última coluna da tabela (%s2)' % ultima,
        isinstance(formula, str) and formula.upper().startswith('=SUM'),
        repr(formula))
    add('O SUM tem valor em cache, senão o cliente abre a célula vazia',
        isinstance(cache, (int, float)), repr(cache))

    # ---- janela ----------------------------------------------------------
    if 'xl/workbook.xml' in partes:
        wbxml = z.read('xl/workbook.xml').decode('utf-8')
        larg = re.search(r'<workbookView[^>]*?\bwindowWidth="(\d+)"', wbxml)
        alt = re.search(r'<workbookView[^>]*?\bwindowHeight="(\d+)"', wbxml)
        atual = (int(larg.group(1)) if larg else 0, int(alt.group(1)) if alt else 0)
        # Limite, não igualdade: o recalc grava JANELA, mas assim que o Daniel
        # abre e salva, o Excel regrava o tamanho real da janela dele (visto:
        # 37080x21000). O que precisa reprovar é o default do LibreOffice,
        # 16384x8192, que abre a planilha em meia tela. 28800 twips = 1920 px.
        add('A planilha abre maximizada no Excel',
            atual[0] >= 28800 and atual[1] >= 15000,
            '%dx%d, mínimo 28800x15000 (o recalc grava %dx%d)'
            % (atual[0], atual[1], JANELA['windowWidth'], JANELA['windowHeight']))

    return checks


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('arquivo')
    ap.add_argument('--itens', type=int, help='Número esperado de demandas')
    ap.add_argument('--total', type=float, help='Total esperado, para bater com o ClickUp')
    ap.add_argument('--nome-entrega', help='Nome do cliente na entrega, para conferir o arquivo')
    args = ap.parse_args()

    checks = conferir(args.arquivo, args.itens, args.total, args.nome_entrega)
    falhas = [c['text'] + ' (' + c['evidence'] + ')' for c in checks if not c['passed']]
    print(json.dumps({'ok': not falhas, 'arquivo': args.arquivo,
                      'falhas': falhas, 'verificacoes': checks},
                     ensure_ascii=False, indent=2))
    sys.exit(0 if not falhas else 2)


if __name__ == '__main__':
    main()
